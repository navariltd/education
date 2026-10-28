# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and contributors
# For license information, please see license.txt

import csv
from io import BytesIO, StringIO

import frappe
import openpyxl
from frappe import _
from frappe.model.document import Document
from frappe.utils import flt, now_datetime
from openpyxl.styles import Alignment, Font

from education.education.doctype.fee_request.fee_request import (
    get_exportable_amount,
    is_term_payable,
    refresh_fee_request,
)
from education.education.doctype.scholar.scholar import get_eligible_scholar_statuses


class FeeRequestExport(Document):
    @frappe.whitelist()
    def fetch_fee_requests(self):
        if (
            self.request_type != "Stipend"
            and self.bank == "KCB"
            and not self.bank_account
        ):
            frappe.throw("Please select a bank account for KCB.")

        FR = frappe.qb.DocType("Fee Request")
        FRT = frappe.qb.DocType("Fee Request Term")
        BA = frappe.qb.DocType("Bank Account")
        SC = frappe.qb.DocType("Scholar")
        SG = frappe.qb.DocType("Scholar Guardian")
        CC = frappe.qb.DocType("County Coordinator Mapping Item")

        exportable = FRT.outstanding_amount - FRT.in_flight_amount
        query = (
            frappe.qb.from_(FR)
            .join(FRT)
            .on((FRT.parent == FR.name) & (FRT.parenttype == "Fee Request"))
            .join(SC)
            .on(FR.scholar == SC.name)
            .left_join(SG)
            .on(
                (SG.parent == SC.name)
                & (SG.parenttype == "Scholar")
                & (SG.parentfield == "guardians")
                & (SG.idx == 1)
            )
            .left_join(BA)
            .on(FR.official_school_name == BA.party)
            .left_join(CC)
            .on(SC.county == CC.county)
            .select(
                FR.name,
                SC.county,
                SG.guardian_name,
                SG.guardian_contact,
                FR.official_school_name,
                FR.scholar,
                FR.student_name,
                exportable.as_("outstanding_amount"),
                BA.custom_branch_name,
                BA.branch_code,
                BA.custom_bank_code,
                BA.bank_account_no,
                BA.bank,
                BA.account_name,
                CC.user,
            )
            .where(
                (FR.docstatus == 1)
                & (FR.academic_year == self.academic_year)
                & (FRT.academic_term == self.academic_term)
                & (FR.company == self.company)
                & (FR.request_type == (self.request_type or "Fees"))
                & (exportable > 0)
                # Requests with Purchase Invoices are paid with Payment Entries, not exports.
                & (FRT.purchase_invoice.isnull() | (FRT.purchase_invoice == ""))
            )
        )

        # A scholar who is no longer eligible is only paid for terms that started on or
        # before their Date of Inactive.
        eligible_statuses = get_eligible_scholar_statuses()
        if eligible_statuses:
            query = query.where(
                SC.status.isin(eligible_statuses)
                | (
                    SC.date_of_inactive.isnotnull()
                    & FRT.term_start_date.isnotnull()
                    & (FRT.term_start_date <= SC.date_of_inactive)
                )
            )

        # A school with several bank accounts, or a county with several coordinators,
        # returns the same request more than once; export it once.
        fee_requests = list(
            {row.name: row for row in reversed(query.run(as_dict=True))}.values()
        )[::-1]

        if not fee_requests:
            frappe.throw("No fee requests found.")

        emails = frappe.db.get_value(
            "County Coordinator Mapping",
            "County Coordinator Mapping",
            ["payable_email", "scholarship_email"],
            as_dict=True,
        )

        results = []
        if self.request_type == "Stipend":
            for fee_request in fee_requests:
                fee_request_details = {
                    "fee_request": fee_request.name,
                    "scholar": fee_request.scholar,
                    "student_name": fee_request.student_name,
                    "school_name": fee_request.official_school_name,
                    "guardian_name": fee_request.guardian_name,
                    "guardian_contact": fee_request.guardian_contact,
                    "amount": fee_request.outstanding_amount,
                }
                results.append(fee_request_details)
        else:
            if self.bank == "Standard Chartered":
                for fee_request in fee_requests:
                    fee_request_details = {
                        "fee_request": fee_request.name,
                        "scholar": fee_request.scholar,
                        "student_name": fee_request.student_name,
                        "school_name": fee_request.official_school_name,
                        "account_number": fee_request.bank_account_no,
                        "bank_code": fee_request.custom_bank_code,
                        "branch_code": fee_request.branch_code,
                        "amount": fee_request.outstanding_amount,
                        "email_address": f"{emails.payable_email},{emails.scholarship_email},{fee_request.user}",
                    }
                    results.append(fee_request_details)

            if self.bank == "KCB":
                company_account = frappe.db.get_value(
                    "Bank Account",
                    {"name": self.bank_account},
                    ["bank_account_no", "branch_code"],
                    as_dict=True,
                )

                for fee_request in fee_requests:
                    fee_request_details = {
                        "fee_request": fee_request.name,
                        "scholar": fee_request.scholar,
                        "student_name": fee_request.student_name,
                        "debit_account": (
                            company_account.get("bank_account_no")
                            if company_account
                            else None
                        ),
                        "beneficiary_name": fee_request.official_school_name,
                        "bank": fee_request.bank,
                        "branch": fee_request.custom_branch_name,
                        "branch_bicsort_code": (
                            company_account.get("branch_code")
                            if company_account
                            else None
                        ),
                        "bicsort_code": fee_request.branch_code,
                        "account_number": fee_request.bank_account_no,
                        "my_reference": fee_request.official_school_name,
                        "sms_notification": fee_request.guardian_contact,
                        "amount": fee_request.outstanding_amount,
                        "email_notification": f"{emails.scholarship_email}",
                    }
                    results.append(fee_request_details)

        return results

    def get_lines(self):
        if self.request_type == "Stipend":
            return self.stipend_requests
        if self.bank == "KCB":
            return self.kcb_fee_requests
        if self.bank == "Standard Chartered":
            return self.standard_chartered_fee_requests
        return []

    def validate(self):
        if self.request_type == "Stipend":
            self.bank = ""

        for row in self.get_lines():
            reference = f"{row.fee_request}|{row.scholar}|{row.student_name}"
            if self.request_type == "Stipend":
                row.reference = reference
            elif self.bank == "Standard Chartered":
                row.details = reference
            else:
                row.beneficiary_reference = reference

    def before_submit(self):
        eligible_statuses = get_eligible_scholar_statuses()
        seen = set()
        for row in self.get_lines():
            if row.fee_request in seen:
                frappe.throw(
                    _("Row {0}: Fee Request {1} is listed more than once.").format(
                        row.idx, row.fee_request
                    )
                )
            seen.add(row.fee_request)

            request_type = frappe.db.get_value(
                "Fee Request", row.fee_request, "request_type"
            )
            if request_type != (self.request_type or "Fees"):
                frappe.throw(
                    _("Row {0}: Fee Request {1} is a {2}, not a {3}.").format(
                        row.idx, row.fee_request, request_type, self.request_type
                    )
                )

            if frappe.db.exists(
                "Fee Request Term",
                {
                    "parent": row.fee_request,
                    "parenttype": "Fee Request",
                    "purchase_invoice": ["is", "set"],
                },
            ):
                frappe.throw(
                    _(
                        "Row {0}: Fee Request {1} uses accounting. Pay its Purchase Invoices with a Payment Entry instead of exporting it."
                    ).format(row.idx, row.fee_request)
                )

            scholar = frappe.db.get_value(
                "Scholar",
                row.scholar
                or frappe.db.get_value("Fee Request", row.fee_request, "scholar"),
                ["name", "status", "date_of_inactive"],
                as_dict=True,
            )
            term_start_date = frappe.db.get_value(
                "Academic Term", self.academic_term, "term_start_date"
            )
            if scholar and not is_term_payable(
                scholar.status,
                scholar.date_of_inactive,
                term_start_date,
                eligible_statuses,
            ):
                frappe.throw(
                    _(
                        "Row {0}: {1} is {2} and {3} starts after their Date of Inactive, so it cannot be exported."
                    ).format(row.idx, scholar.name, scholar.status, self.academic_term)
                    if scholar.date_of_inactive
                    else _(
                        "Row {0}: {1} is {2} with no Date of Inactive, so it cannot be exported."
                    ).format(row.idx, scholar.name, scholar.status)
                )

            exportable = get_exportable_amount(row.fee_request, self.academic_term)
            if flt(row.amount) <= 0 or flt(row.amount) > flt(exportable) + 0.005:
                frappe.throw(
                    _(
                        "Row {0}: {1} can export at most {2} for {3}. Get Fee Requests again."
                    ).format(
                        row.idx,
                        row.fee_request,
                        max(flt(exportable), 0),
                        self.academic_term,
                    )
                )

            row.status = "Pending"
            row.paid_amount = 0
            row.settled_by_doctype = row.settled_by = row.withdraw_reason = None

    def on_submit(self):
        self.db_set("submitted_on", now_datetime())
        self.refresh_fee_requests()

    def before_cancel(self):
        if any(row.status == "Settled" for row in self.get_lines()):
            frappe.throw(
                _(
                    "Some lines have already been paid, so this export cannot be cancelled. Withdraw the unpaid lines instead."
                )
            )

    def on_cancel(self):
        self.ignore_linked_doctypes = ("Fee Request", "Fee Request Payment Import")
        self.refresh_fee_requests()

    def refresh_fee_requests(self, fee_requests=None):
        for name in sorted(
            fee_requests or {row.fee_request for row in self.get_lines()}
        ):
            if name:
                refresh_fee_request(name)

    @frappe.whitelist()
    def withdraw_lines(self, lines, reason):
        if self.docstatus != 1:
            frappe.throw(_("Only lines of a submitted export can be withdrawn."))

        self.check_permission("submit")

        if not (reason or "").strip():
            frappe.throw(_("A reason is required."))

        lines = set(frappe.parse_json(lines) or [])
        withdrawn = set()
        for row in self.get_lines():
            if row.name not in lines:
                continue

            if row.status != "Pending":
                frappe.throw(
                    _("Row {0} is {1}. Only Pending lines can be withdrawn.").format(
                        row.idx, row.status
                    )
                )

            row.db_set({"status": "Withdrawn", "withdraw_reason": reason.strip()})
            withdrawn.add(row.fee_request)

        self.refresh_fee_requests(withdrawn)
        return len(withdrawn)


@frappe.whitelist()
def export_fee_requests(export_docname, format="excel"):
    """Main entry point"""

    doc = frappe.get_doc("Fee Request Export", export_docname)

    if (
        doc.request_type != "Stipend"
        and not doc.kcb_fee_requests
        and not doc.standard_chartered_fee_requests
    ):
        frappe.throw("No fee requests to export")

    headers, data = get_headers_and_data(doc)

    if format.lower() == "excel":
        stream_excel(doc.name, headers, data)
    else:
        stream_csv(doc.name, headers, data)


def get_headers_and_data(doc):
    headers = []
    data = []

    if doc.request_type == "Stipend":
        headers = [
            "STUDENT NAME",
            "REFERENCE",
            "OFFICIAL SCHOOL NAME",
            "GUARDIAN NAME",
            "GUARDIAN CONTACT",
            "AMOUNT",
        ]

        data = [
            [
                row.student_name,
                row.reference,
                row.official_school_name,
                row.guardian_name,
                row.guardian_contact,
                row.amount,
            ]
            for row in doc.stipend_requests
            if row.status != "Withdrawn"
        ]
    else:
        if doc.bank == "Standard Chartered":
            headers = [
                "NAME",
                "ACCOUNT NO",
                "BANK CODE",
                "BRANCH CODE",
                "AMOUNT",
                "EMAIL ADDRESS",
                "DETAILS",
            ]

            data = [
                [
                    row.student_name,
                    row.account_number,
                    row.bank_code,
                    row.branch_code,
                    row.amount,
                    row.email_address,
                    row.details,
                ]
                for row in doc.standard_chartered_fee_requests
                if row.status != "Withdrawn"
            ]

        elif doc.bank == "KCB":
            headers = [
                "Debit Account",
                "Branch BIC/SORT Code",
                "Beneficiary Name",
                "Bank",
                "Branch",
                "BIC/SORT Code",
                "Account Number",
                "My Reference",
                "Beneficiary Reference",
                "Amount",
                "SMS Notification",
                "Email Notification",
            ]

            data = [
                [
                    row.debit_account,
                    row.branch_bicsort_code,
                    row.beneficiary_name,
                    row.bank,
                    row.branch,
                    row.bicsort_code,
                    row.account_number,
                    row.my_reference,
                    row.beneficiary_reference,
                    row.amount,
                    row.sms_notification,
                    row.email_notification,
                ]
                for row in doc.kcb_fee_requests
                if row.status != "Withdrawn"
            ]

        else:
            frappe.throw("Unsupported bank type")

    return headers, data


def stream_excel(name, headers, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fee Requests"

    # Metadata
    ws.append(["Fee Request Export", name])
    ws.append(["Export Date:", frappe.utils.getdate()])
    ws.append([])

    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")

    # Headers
    ws.append(headers)

    header_row = ws.max_row
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.font = header_font
        cell.alignment = header_alignment

    for row in data:
        ws.append(row)

    # Auto column width
    for column in ws.columns:
        max_length = 0
        col_letter = column[0].column_letter

        for cell in column:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception as e:  # E722
                frappe.log_error(title="Error while adjusting column width", message=e)

        ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    frappe.response.filename = f"{name}_fee_requests.xlsx"
    frappe.response.filecontent = output.getvalue()
    frappe.response.type = "binary"


def stream_csv(name, headers, data):
    output = StringIO()
    writer = csv.writer(output)

    # Metadata
    writer.writerow(["Fee Request Export:", name])
    writer.writerow(["Export Date:", frappe.utils.getdate()])
    writer.writerow([])

    # Headers
    writer.writerow(headers)

    # Data
    writer.writerows(data)

    frappe.response.filename = f"{name}_fee_requests.csv"
    frappe.response.filecontent = output.getvalue()
    frappe.response.type = "csv"
